import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

EXCEL_FILE = "eventos.xlsx"

COLUMNS = [
    "Fecha del Evento",
    "Hora del Evento",
    "Tipo de RED",
    "Site",
    "Categoría",
    "Título del Evento",
    "Determinante y Tienda",
    "Descripción del Evento",
    "Planes de Acción o seguimiento",
    "Personal Enterado",
    "Importe",
    "Fecha de Registro",
]

KEY_MAP = {
    "Fecha del Evento":               "fecha_evento",
    "Hora del Evento":                "hora_evento",
    "Tipo de RED":                    "tipo_red",
    "Site":                           "site",
    "Categoría":                      "categoria",
    "Título del Evento":              "titulo_evento",
    "Determinante y Tienda":          "determinante_tienda",
    "Descripción del Evento":         "descripcion_evento",
    "Planes de Acción o seguimiento": "planes_accion",
    "Personal Enterado":              "personal_enterado",
    "Importe":                        "importe",
}

HEADER_FILL  = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT  = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _build_workbook() -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Eventos"
        for col_idx, header in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font  = HEADER_FONT
            cell.fill  = HEADER_FILL
            cell.alignment = HEADER_ALIGN
        ws.row_dimensions[1].height = 35
    return wb, ws


def _adjust_column_widths(ws) -> None:
    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max(
            (len(str(cell.value)) if cell.value else 0 for cell in col),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def save_event(event_data: dict) -> str:
    wb, ws = _build_workbook()

    row = []
    for col in COLUMNS:
        if col == "Fecha de Registro":
            row.append(datetime.now().strftime("%d/%m/%Y %H:%M:%S"))
        else:
            row.append(event_data.get(KEY_MAP[col]) or "")

    ws.append(row)

    last_row = ws.max_row
    for cell in ws[last_row]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    _adjust_column_widths(ws)
    wb.save(EXCEL_FILE)
    return os.path.abspath(EXCEL_FILE)
